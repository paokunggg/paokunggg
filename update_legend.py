#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""สร้าง legend.json จากไฟล์ Excel แล้ว commit + push ขึ้น GitHub

ใช้ชีต "legend from typical" คอลัมน์:
  A = LegendA (รหัสหลัก)  B = Legend (รหัสย่อย)
  C = Detail              D = Quantity    E = material No.

ผลลัพธ์: [[legendA, [[detail, qty, matNo, sub], ...]], ...]
"""
import json
import os
import subprocess
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SHEET = 'legend from typical'
OUT = os.path.join(HERE, 'legend.json')

# ลองไฟล์ตามลำดับ ใช้ไฟล์แรกที่เจอ
SOURCES = ['AR.xlsx', 'legend ที่รวมรวมมา.xlsx']


def find_source():
    for name in SOURCES:
        path = os.path.join(HERE, name)
        if os.path.exists(path):
            return path
    sys.exit('ไม่พบไฟล์ Excel: ' + ' หรือ '.join(SOURCES))


def num(v):
    """แปลงจำนวนเป็น float ถ้าทำได้ ไม่ได้ก็คืนค่าเดิมเป็นข้อความ"""
    if v is None or v == '':
        return 0
    try:
        return float(str(v).strip().replace(',', ''))
    except ValueError:
        return str(v).strip()


def text(v):
    return '' if v is None else str(v).strip()


def build(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit('ไม่พบชีต "%s" ในไฟล์ %s' % (SHEET, os.path.basename(path)))
    ws = wb[SHEET]

    order = []          # เก็บลำดับ legendA ตามที่ปรากฏในไฟล์
    groups = {}
    skipped = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        legend_a = text(row[0])
        if not legend_a:
            skipped += 1
            continue
        sub = text(row[1])
        detail = text(row[2])
        qty = num(row[3])
        mat = text(row[4])

        if legend_a not in groups:
            groups[legend_a] = []
            order.append(legend_a)
        groups[legend_a].append([detail, qty, mat, sub])

    wb.close()
    data = [[code, groups[code]] for code in order]
    return data, skipped


def git(*args):
    return subprocess.run(['git'] + list(args), cwd=HERE,
                          capture_output=True, text=True, encoding='utf-8')


def main():
    src = find_source()
    print('อ่านข้อมูลจาก: %s' % os.path.basename(src))

    data, skipped = build(src)
    items = sum(len(g[1]) for g in data)
    print('พบ Legend %d รหัส  รวม %d รายการ  (ข้ามแถวว่าง %d แถว)'
          % (len(data), items, skipped))

    if not data:
        sys.exit('ไม่มีข้อมูล — ยกเลิก ไม่เขียนทับ legend.json')

    old = None
    if os.path.exists(OUT):
        with open(OUT, encoding='utf-8') as f:
            old = f.read()

    new = json.dumps(data, ensure_ascii=False)
    if old == new:
        print('ข้อมูลเหมือนเดิม ไม่มีอะไรต้องอัปเดต')
        return

    with open(OUT, 'w', encoding='utf-8') as f:
        f.write(new)
    print('เขียน legend.json เรียบร้อย (%.0f KB)' % (len(new.encode('utf-8')) / 1024))

    # commit + push
    if git('rev-parse', '--git-dir').returncode != 0:
        print('ไม่ได้อยู่ใน git repository — ข้ามการ push')
        return

    git('add', 'legend.json')
    res = git('commit', '-m', 'update legend data')
    if res.returncode != 0:
        print('ไม่มีการเปลี่ยนแปลงให้ commit')
        return
    print('commit เรียบร้อย')

    res = git('push')
    if res.returncode != 0:
        print('push ไม่สำเร็จ:')
        print(res.stderr.strip())
        sys.exit(1)
    print('push ขึ้น GitHub เรียบร้อย — รอ 1-2 นาที แล้วกด Ctrl+F5 ที่หน้าเว็บ')


if __name__ == '__main__':
    main()
